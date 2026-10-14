import io
from pathlib import Path
from unittest.mock import MagicMock, patch

import numpy as np
import pytest
from PIL import Image

from app import app
from pipeline.document_reconstructor import PageLayout, LayoutElement


def _make_dummy_image_bytes():
    img = Image.new("RGB", (100, 100), color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


def test_convert_rejects_missing_file(client):
    res = client.post("/convert", data={"file_name": "test.png"})
    assert res.status_code == 400
    assert "No file uploaded" in res.get_json()["error"]


def test_convert_rejects_unsupported_format(client):
    res = client.post(
        "/convert",
        data={"file_name": "malicious.exe", "file": (io.BytesIO(b"executable"), "malicious.exe")},
    )
    assert res.status_code == 400
    assert "Unsupported file format" in res.get_json()["error"]


def test_download_prevents_path_traversal(client):
    # Attempt directory traversal attacks
    res1 = client.get("/download/../../app.py")
    assert res1.status_code in {400, 404}

    res2 = client.get("/download/..%2F..%2Fetc%2Fpasswd")
    assert res2.status_code in {400, 404}

    res3 = client.get("/download/non_existent_document.docx")
    assert res3.status_code == 404


@patch("app.get_engine")
def test_convert_multi_page_pdf_success(mock_get_engine, client, tmp_path):
    mock_engine = MagicMock()
    mock_get_engine.return_value = mock_engine

    dummy_docx = tmp_path / "test_output.docx"
    dummy_docx.write_bytes(b"PK\x03\x04dummy-docx-content")

    # Simulate 3-page layout result
    dummy_layouts = [
        PageLayout(page_index=0, width=595, height=842, elements=[]),
        PageLayout(page_index=1, width=595, height=842, elements=[]),
        PageLayout(page_index=2, width=595, height=842, elements=[]),
    ]
    mock_engine.process_file_bytes.return_value = (dummy_docx, dummy_layouts)

    # Valid minimal PDF header
    fake_pdf_bytes = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"

    res = client.post(
        "/convert",
        data={"file_name": "contracts_multi_page.pdf", "file": (io.BytesIO(fake_pdf_bytes), "contracts_multi_page.pdf")},
    )

    assert res.status_code == 200
    data = res.get_json()
    assert data["message"] == "OCR process completed successfully"
    assert data["total_pages"] == 3
    assert data["docx_file"].startswith("contracts_multi_page_")
    assert data["docx_file"].endswith(".docx")
    assert data["download_url"].startswith("/download/")


def test_pipeline_integration_real_pdf_to_docx(tmp_path):
    """Pipeline integration test:
    Creates a real 2-page PDF, runs process_file_bytes() with simulated OCR layout,
    and verifies that the resulting DOCX contains paragraphs, table contents, and correct page sections.
    """
    import fitz
    from docx import Document
    from unittest.mock import patch
    from pipeline.unified_engine import UnifiedDocumentEngine
    from pipeline.adaptive_table_extractor import TableResult, TableCell


    # 1. Create a real 2-page PDF using PyMuPDF
    doc = fitz.open()
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text((50, 100), "Trang 1: Hop dong thuong mai", fontsize=14)
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((50, 100), "Trang 2: Bang gia chi tiet", fontsize=14)
    pdf_bytes = doc.tobytes()
    doc.close()

    # 2. Mock reconstructor with simulated process_page to avoid initializing PaddleOCR
    mock_reconstructor = MagicMock()

    def mock_process_page(img_bgr, page_index=0, dpi=150):
        if page_index == 0:
            elem = LayoutElement(
                page_index=0,
                element_type="paragraph",
                bbox=(50, 50, 500, 80),
                content="Cong Hoa Xa Hoi Chu Nghia Viet Nam",
                metadata={
                    "font_name": "Calibri",
                    "font_size_pt": 14,
                    "bold": True,
                    "alignment": "center",
                },
            )
            return PageLayout(page_index=0, width=595, height=842, elements=[elem])
        else:
            table_res = TableResult(
                rows=[
                    [TableCell(row=0, col=0, text="STT"), TableCell(row=0, col=1, text="Hang Hoa")],
                    [TableCell(row=1, col=0, text="1"), TableCell(row=1, col=1, text="Camera HIKVISION")],
                ],
                spans=[],
                column_widths_pct=[0.2, 0.8],
            )
            elem = LayoutElement(
                page_index=1,
                element_type="table",
                bbox=(50, 100, 400, 200),
                content=table_res,
            )
            return PageLayout(page_index=1, width=595, height=842, elements=[elem])

    mock_reconstructor.process_page.side_effect = mock_process_page

    # Inject mock_reconstructor directly - zero heavy OCR initialization
    engine = UnifiedDocumentEngine(
        artifacts_dir=str(tmp_path / "artifacts"),
        reconstructor=mock_reconstructor,
    )
    out_docx_path = tmp_path / "integrated_output.docx"

    saved_docx, layouts = engine.process_file_bytes(
        file_bytes=pdf_bytes,
        file_kind="pdf",
        output_docx_path=out_docx_path,
        dpi=72,
    )


    # 3. Verify execution and file generation
    assert saved_docx.exists()
    assert len(layouts) == 2

    # 4. Verify generated Word document structure and contents
    word_doc = Document(str(saved_docx))
    assert len(word_doc.sections) == 2

    # Verify paragraph content from page 1 was reconstructed
    doc_text = " ".join([p.text for p in word_doc.paragraphs])
    assert "Cong Hoa Xa Hoi Chu Nghia Viet Nam" in doc_text

    # Verify table structure from page 2 was reconstructed
    assert len(word_doc.tables) >= 1
    table = word_doc.tables[0]
    cell_texts = [c.text.strip() for row in table.rows for c in row.cells]
    assert "STT" in cell_texts
    assert "Camera HIKVISION" in cell_texts


def test_dashboard_first_page_cached_flow(tmp_path):
    """Verifies that language detection with return_first_page=True passes
    cached_first_page into process_to_layouts without re-rasterizing page 0.
    """
    import fitz
    from unittest.mock import MagicMock
    from pipeline.unified_engine import UnifiedDocumentEngine

    doc = fitz.open()
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text((50, 100), "Trang 1 text", fontsize=14)
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((50, 100), "Trang 2 text", fontsize=14)
    pdf_bytes = doc.tobytes()
    doc.close()

    mock_reconstructor = MagicMock()
    mock_reconstructor.process_page.return_value = PageLayout(page_index=0, width=595, height=842, elements=[])

    mock_lang_detector = MagicMock()
    mock_lang_detector.detect_image_language.return_value = ("vi", 0.99, {})

    engine = UnifiedDocumentEngine(
        artifacts_dir=str(tmp_path / "artifacts"),
        reconstructor=mock_reconstructor,
        lang_detector=mock_lang_detector,
    )

    # 1. Detect language and get cached first page
    lang, conf, details, first_page = engine.detect_file_language(
        pdf_bytes, file_kind="pdf", dpi=72, return_first_page=True
    )
    assert lang == "vi"
    assert isinstance(first_page, np.ndarray)

    # 2. Process to layouts passing cached first page
    layouts = engine.process_to_layouts(
        file_bytes=pdf_bytes,
        file_kind="pdf",
        dpi=72,
        cached_first_page=first_page,
    )
    assert len(layouts) == 2
    assert mock_reconstructor.process_page.call_count == 2
    # Verify the first page passed to process_page was exactly our cached array
    passed_first_img = mock_reconstructor.process_page.call_args_list[0][0][0]
    assert np.array_equal(passed_first_img, first_page)



