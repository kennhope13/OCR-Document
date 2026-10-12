import pytest
from pathlib import Path
from docx import Document
from pipeline.adaptive_table_extractor import TableResult, TableCell
from pipeline.docx_builder import DocxBuilder

def test_docx_builder_simple_table(tmp_path):
    builder = DocxBuilder()
    doc = builder.create_document()

    cells = [
        [TableCell(row=0, col=0, text="Model"), TableCell(row=0, col=1, text="Qty")],
        [TableCell(row=1, col=0, text="DS-123"), TableCell(row=1, col=1, text="10PCS")]
    ]
    t_result = TableResult(
        rows=cells,
        spans=[],
        column_widths_pct=[70.0, 30.0],
        extraction_method="test"
    )

    builder.add_table(doc, t_result)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.rows) == 2
    assert len(t.columns) == 2
    assert t.cell(0, 0).text == "Model"
    assert t.cell(1, 0).text == "DS-123"
    assert t.cell(1, 1).text == "10PCS"

def test_docx_builder_with_merged_cells(tmp_path):
    builder = DocxBuilder()
    doc = builder.create_document()

    # Bảng 2x2 với hàng đầu merge 2 cột
    cells = [
        [TableCell(row=0, col=0, text="Header Span", colspan=2), TableCell(row=0, col=1, is_spanned_copy=True)],
        [TableCell(row=1, col=0, text="A"), TableCell(row=1, col=1, text="B")]
    ]
    spans = [{"row": 0, "col": 0, "rowspan": 1, "colspan": 2}]
    t_result = TableResult(
        rows=cells,
        spans=spans,
        column_widths_pct=[50.0, 50.0],
        extraction_method="test"
    )

    builder.add_table(doc, t_result)
    t = doc.tables[0]
    # Khi merge ô (0,0) và (0,1), text ở cả 2 ô trỏ chung về 1 vùng
    assert t.cell(0, 0).text == "Header Span"
    assert t.cell(0, 1).text == "Header Span"
    assert t.cell(1, 0).text == "A"
    assert t.cell(1, 1).text == "B"

    out_file = tmp_path / "test_out.docx"
    builder.build([t_result], out_file)
    assert out_file.exists()
