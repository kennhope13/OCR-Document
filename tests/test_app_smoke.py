import io
import os
from unittest.mock import MagicMock, patch

import numpy as np
from docx import Document
from PIL import Image

from app import app
from ocr.ocr_handler import process_image, process_pdf


def _make_png_bytes():
    image = Image.new('RGB', (20, 20), color='white')
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    return buffer.getvalue()


def test_flask_routes_exist():
    client = app.test_client()

    res_convert = client.get('/convert')
    res_download = client.get('/download/not-found.docx')

    assert res_convert.status_code in {400, 405, 404}
    assert res_download.status_code in {404, 405}


@patch('app.storage.Client')
@patch('app.create_docx')
@patch('app.process_image')
@patch('app.deskew')
def test_convert_image_success_with_mocked_dependencies(mock_deskew, mock_process_image, mock_create_docx, mock_client):
    client = app.test_client()
    png_bytes = _make_png_bytes()

    mock_deskew.return_value = np.array(Image.new('RGB', (20, 20), color='white'))
    mock_process_image.return_value = [[[[[0, 0], [10, 0], [10, 10], [0, 10]], ('Hello OCR', 0.99)]]]

    fake_blob = MagicMock()
    fake_blob.upload_from_filename = MagicMock()
    fake_bucket = MagicMock()
    fake_bucket.blob.return_value = fake_blob
    mock_client.return_value.bucket.return_value = fake_bucket

    with open('tmp_test_output.docx', 'wb') as f:
        f.write(b'fake-docx')
    mock_create_docx.return_value = 'tmp_test_output.docx'

    response = client.post('/convert', data={'file_name': 'sample.png', 'file': (io.BytesIO(png_bytes), 'sample.png')})

    assert response.status_code == 200
    data = response.get_json()
    assert data['message'] == 'OCR process completed'
    assert 'docx_file' in data
    assert data['download_url'].startswith('/download/')

    if os.path.exists('tmp_test_output.docx'):
        os.remove('tmp_test_output.docx')


@patch('app.storage.Client')
@patch('app.process_pdf')
@patch('app.detect_file_kind')
@patch('app.is_supported_file')
def test_convert_invalid_pdf_returns_error(mock_supported, mock_kind, mock_pdf, mock_client):
    client = app.test_client()

    mock_supported.return_value = True
    mock_kind.return_value = 'pdf'
    mock_pdf.side_effect = RuntimeError('PDF processing failed')

    response = client.post('/convert', data={'file_name': 'bad.pdf', 'file': (io.BytesIO(b'%PDF-1.4'), 'bad.pdf')})

    assert response.status_code == 500
    assert 'error' in response.get_json()


@patch('app.storage.Client')
def test_download_uses_gcs_when_local_file_missing(mock_client):
    client = app.test_client()

    fake_blob = MagicMock()
    fake_blob.exists.return_value = True
    fake_blob.download_as_bytes.return_value = b'fake-docx-content'
    fake_bucket = MagicMock()
    fake_bucket.blob.return_value = fake_blob
    mock_client.return_value.bucket.return_value = fake_bucket

    response = client.get('/download/generated.docx')

    assert response.status_code == 200
    assert response.data == b'fake-docx-content'


@patch('app.storage.Client')
def test_download_missing_file_returns_404(mock_client):
    client = app.test_client()

    fake_blob = MagicMock()
    fake_blob.exists.return_value = False
    fake_bucket = MagicMock()
    fake_bucket.blob.return_value = fake_blob
    mock_client.return_value.bucket.return_value = fake_bucket

    response = client.get('/download/missing.docx')

    assert response.status_code == 404
    assert response.get_json()['error'] == 'File not found'


def _ocr_result_for_text(text):
    coords = [[0, 0], [200, 0], [200, 50], [0, 50]]
    return [[[coords, (text, 0.99)]]]


def _make_jpg_bytes():
    image = Image.new('RGB', (80, 40), color='white')
    buffer = io.BytesIO()
    image.save(buffer, format='JPEG')
    return buffer.getvalue()


def _make_png_bytes_for_integration():
    image = Image.new('RGB', (80, 40), color='white')
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    return buffer.getvalue()


def _make_valid_pdf_bytes():
    return b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n4 0 obj\n<< /Length 44 >>\nstream\nBT /F1 18 Tf 50 80 Td (Quarterly report) Tj ET\nendstream\nendobj\n5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\nxref\n0 6\n0000000000 65535 f \n0000000010 00000 n \n0000000062 00000 n \n0000000123 00000 n \n0000000245 00000 n \n0000000567 00000 n \ntrailer\n<< /Root 1 0 R /Size 6 >>\nstartxref\n700\n%%EOF'


@patch('ocr.ocr_handler.get_ocr_model')
def test_process_image_creates_docx_for_valid_jpg_and_png(mock_get_ocr_model, monkeypatch, tmp_path):
    mock_get_ocr_model.return_value.ocr.return_value = _ocr_result_for_text('Invoice 2026')

    for file_name, file_bytes in [('sample.jpg', _make_jpg_bytes()), ('sample.png', _make_png_bytes_for_integration())]:
        monkeypatch.chdir(tmp_path)
        process_image(io.BytesIO(file_bytes), file_name)
        doc_path = tmp_path / f'{file_name.lower().replace(".jpg", ".docx").replace(".png", ".docx")}'
        doc = Document(doc_path)
        text = ' '.join(paragraph.text for paragraph in doc.paragraphs)
        assert 'Invoice 2026' in text
        doc_path.unlink()


@patch('ocr.ocr_handler.convert_from_path')
@patch('ocr.ocr_handler.get_ocr_model')
def test_process_pdf_creates_docx_for_valid_pdf(mock_get_ocr_model, mock_convert_from_path, monkeypatch, tmp_path):
    mock_get_ocr_model.return_value.ocr.return_value = _ocr_result_for_text('Quarterly report')
    mock_convert_from_path.return_value = [Image.new('RGB', (80, 40), color='white')]

    monkeypatch.chdir(tmp_path)
    process_pdf(_make_valid_pdf_bytes(), file_name='sample.pdf')

    doc_path = tmp_path / 'sample_page_1.docx'
    doc = Document(doc_path)
    text = ' '.join(paragraph.text for paragraph in doc.paragraphs)
    assert 'Quarterly report' in text
    doc_path.unlink()


@patch('app.deskew')
@patch('app.process_image')
@patch('app.storage.Client')
@patch('app.create_docx')
def test_app_convert_accepts_real_jpg_upload_and_writes_docx(mock_create_docx, mock_client, mock_process_image, mock_deskew, tmp_path):
    client = app.test_client()
    image_bytes = _make_jpg_bytes()
    fake_blob = MagicMock()
    fake_blob.upload_from_filename = MagicMock()
    fake_bucket = MagicMock()
    fake_bucket.blob.return_value = fake_blob
    mock_client.return_value.bucket.return_value = fake_bucket
    mock_deskew.return_value = np.array(Image.new('RGB', (80, 40), color='white'))
    mock_process_image.return_value = _ocr_result_for_text('Invoice 2026')

    mock_create_docx.return_value = str(tmp_path / 'real-output.docx')
    with open(str(tmp_path / 'real-output.docx'), 'wb') as handle:
        handle.write(b'fake-docx')

    response = client.post(
        '/convert',
        data={'file_name': 'sample.jpg', 'file': (io.BytesIO(image_bytes), 'sample.jpg')},
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload['message'] == 'OCR process completed'
    assert payload['download_url'].endswith('/download/real-output.docx')


@patch('app.storage.Client')
@patch('app.process_image', return_value=None)
def test_app_convert_returns_clear_error_for_empty_ocr_result(mock_process_image, mock_client):
    client = app.test_client()
    mock_client.return_value.bucket.return_value = MagicMock()

    response = client.post(
        '/convert',
        data={'file_name': 'sample.jpg', 'file': (io.BytesIO(_make_jpg_bytes()), 'sample.jpg')},
    )

    assert response.status_code == 400
    assert response.get_json()['error'] == 'OCR result is empty or invalid.'
