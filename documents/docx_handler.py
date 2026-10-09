try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.section import WD_ORIENT
except ImportError:  # pragma: no cover - optional dependency for tests
    Document = None
    Inches = Pt = None
    WD_ORIENT = None


# Tạo docx
def create_docx(ocr_result, file_name):
    if Document is None:
        raise RuntimeError("python-docx is not installed. Please install project dependencies.")

    min_width = float('inf')
    doc = Document()
    for item in ocr_result[0]:
        coordinates = item[0]
        text = item[1]
        text_values = text[0]
        x1, y1 = coordinates[0][0], coordinates[0][1]

        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text_values)

        width = coordinates[1][0] - coordinates[0][0]
        height = coordinates[2][1] - coordinates[0][1]
        font_size = min(width, height)
        run.font.size = Pt(font_size)

        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(x1)
        paragraph.paragraph_format.space_before = Pt(y1 * 0.000001)
        if x1 < min_width:
            min_width = x1

        section = doc.sections[-1]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(min_width)

    output_path = file_name.lower().replace('.pdf', '.docx').replace('.jpg', '.docx').replace('.jpeg', '.docx').replace('.png', '.docx')
    if not output_path.endswith('.docx'):
        output_path = f"{output_path}.docx"

    doc.save(output_path)
    return output_path

