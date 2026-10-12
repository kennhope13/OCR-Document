from pathlib import Path
from typing import List, Union, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pipeline.adaptive_table_extractor import TableResult, TableCell
from pipeline.document_reconstructor import PageLayout, LayoutElement

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))


class DocxBuilder:
    def __init__(self, page_width_in: float = 8.27, page_height_in: float = 11.69, margin_in: float = 0.5):
        self.page_width_in = page_width_in
        self.page_height_in = page_height_in
        self.margin_in = margin_in
        self.printable_width_in = max(1.0, page_width_in - 2 * margin_in)

    def create_document(self) -> Document:
        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(self.page_width_in)
        section.page_height = Inches(self.page_height_in)
        section.top_margin = Inches(self.margin_in)
        section.bottom_margin = Inches(self.margin_in)
        section.left_margin = Inches(self.margin_in)
        section.right_margin = Inches(self.margin_in)

        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Calibri'
        normal_style.font.size = Pt(8.5)
        normal_style.paragraph_format.line_spacing = 1.05
        normal_style.paragraph_format.space_after = Pt(1)
        return doc

    def add_table(self, document: Document, table_result: TableResult, printable_width_in: Optional[float] = None) -> None:
        if not table_result.rows:
            return

        num_rows = len(table_result.rows)
        num_cols = len(table_result.rows[0])
        if num_cols == 0:
            return

        table = document.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        eff_w = printable_width_in if printable_width_in is not None else self.printable_width_in

        for r_idx, row in enumerate(table_result.rows):
            for c_idx, cell_data in enumerate(row):
                word_cell = table.cell(r_idx, c_idx)
                if not cell_data.is_spanned_copy:
                    word_cell.text = cell_data.text
                    if cell_data.image_path and Path(cell_data.image_path).exists():
                        img_p = word_cell.add_paragraph()
                        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        img_run = img_p.add_run()
                        img_run.add_picture(cell_data.image_path, width=Inches(2.0))

        for span in table_result.spans:
            r = span["row"]
            c = span["col"]
            r_span = span.get("rowspan", 1)
            c_span = span.get("colspan", 1)

            end_r = min(num_rows - 1, r + r_span - 1)
            end_c = min(num_cols - 1, c + c_span - 1)
            if end_r > r or end_c > c:
                start_cell = table.cell(r, c)
                end_cell = table.cell(end_r, end_c)
                start_cell.merge(end_cell)

        pcts = table_result.column_widths_pct
        if len(pcts) == num_cols and sum(pcts) > 0:
            col_widths_inches = [(p / 100.0) * eff_w for p in pcts]
        else:
            col_widths_inches = [eff_w / num_cols for _ in range(num_cols)]

        for row in table.rows:
            for c_i, w_in in enumerate(col_widths_inches):
                row.cells[c_i].width = Inches(w_in)

        if table_result.extraction_method == "header_block_2col":
            for r in table.rows:
                r.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                c1 = r.cells[1]
                c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                raw_txt = c1.text.strip()
                c1.text = ""
                for line in [l.strip() for l in raw_txt.split("\n") if l.strip()]:
                    p = c1.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(line)
                    run.font.name = 'Times New Roman'
                    if "CỘNG HÒA" in line.upper():
                        run.font.size = Pt(12)
                        run.font.bold = True
                    elif "ĐỘC LẬP" in line.upper() or "HẠNH PHÚC" in line.upper():
                        run.font.size = Pt(12)
                        run.font.bold = True
                    elif any(k in line.lower() for k in ["ngày", "tháng", "năm"]):
                        run.font.size = Pt(11)
                        run.font.italic = True
                    else:
                        run.font.size = Pt(11)
            return

        # Kiểm tra xem hàng 0 và hàng 1 có phải là header/subheader không
        has_sub_header = any(s.get("rowspan", 1) > 1 for s in table_result.spans if s.get("row") == 0)

        is_borderless = ("borderless" in table_result.extraction_method) or ("signature" in table_result.extraction_method)

        for r_i, row in enumerate(table.rows):
            for c_i, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if not is_borderless:
                    set_cell_border(
                        cell,
                        top={"sz": 4, "val": "single", "color": "B0B0B0"},
                        bottom={"sz": 4, "val": "single", "color": "B0B0B0"},
                        left={"sz": 4, "val": "single", "color": "B0B0B0"},
                        right={"sz": 4, "val": "single", "color": "B0B0B0"}
                    )
                is_header_row = (r_i == 0) or (has_sub_header and r_i == 1)

                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    p.paragraph_format.line_spacing = 1.1

                    # Canh lề cho ô
                    cell_text = cell.text.strip()
                    if is_borderless or is_header_row or len(cell_text) <= 5 or c_i in [0, 4, 5, 7]:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        if is_borderless:
                            font_size_val = 11.5 if (is_header_row or r_i == len(table.rows) - 1) else 10.5
                            run.font.size = Pt(font_size_val)
                            if is_header_row:
                                run.font.bold = True
                        else:
                            font_size_val = 9.5 if eff_w >= 9.5 else (8.5 if num_cols <= 8 else 7.5)
                            run.font.size = Pt(font_size_val)
                            if is_header_row and not table_result.extraction_method.startswith("form_e_boxes"):
                                run.font.bold = True

    def add_page_layout(self, document: Document, page: PageLayout, printable_width_in: Optional[float] = None) -> None:
        """
        Dựng tuần tự các phần tử trên một trang theo đúng thứ tự reading order đã sắp xếp.
        """
        eff_w = printable_width_in if printable_width_in is not None else self.printable_width_in

        for elem in page.elements:
            if elem.element_type == "table":
                self.add_table(document, elem.content, printable_width_in=eff_w)

            elif elem.element_type == "image":
                img_path = str(elem.content)
                if Path(img_path).exists():
                    p = document.add_paragraph()
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run()
                    if elem.metadata.get("kind") == "stamp_seal":
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        w_px = elem.metadata.get("width_px", 75)
                        stamp_w_in = min(1.8, max(0.4, w_px / 150.0))
                        run.add_picture(img_path, width=Inches(stamp_w_in))
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run.add_picture(img_path, width=Inches(min(eff_w, 6.5)))

            elif elem.element_type in ["paragraph", "header"]:
                p = document.add_paragraph()
                meta = elem.metadata or {}

                # 1. Căn lề
                align = meta.get("alignment", "left")
                if align == "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align == "right":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif align == "justify":
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # 2. Thụt lề (Indentation)
                indent_in = meta.get("indent_left_in", 0.0)
                if indent_in > 0:
                    p.paragraph_format.left_indent = Inches(indent_in)

                # 3. Giãn dòng và khoảng cách trước/sau
                p.paragraph_format.space_before = Pt(meta.get("space_before_pt", 1.5))
                p.paragraph_format.space_after = Pt(meta.get("space_after_pt", 1.5))
                p.paragraph_format.line_spacing = 1.15

                # 4. Phông chữ, Cỡ chữ, Màu sắc & Kiểu chữ
                full_text = str(elem.content)
                color = meta.get("color_rgb")
                field_prefixes = [
                    "Đại diện bên giao:", "Đại diện bên giao",
                    "Đại diện bên nhận:", "Đại diện bên nhận",
                    "Đại diện bên A:", "Đại diện bên A",
                    "Đại diện bên B:", "Đại diện bên B",
                    "Chức vụ:", "Chức vụ", "Chức vi:", "Chức vi",
                    "Địa chỉ:", "Địa chỉ giao hàng:", "Điện thoại:", "Email:", "MST:"
                ]
                has_prefix = any(pfx.lower() in full_text.lower() for pfx in field_prefixes)

                if color and has_prefix:
                    import re
                    pattern = '(' + '|'.join(map(re.escape, field_prefixes)) + ')'
                    tokens = [t for t in re.split(pattern, full_text, flags=re.I) if t]
                    for tok in tokens:
                        r = p.add_run(tok)
                        r.font.name = meta.get("font_name", "Times New Roman")
                        r.font.size = Pt(meta.get("font_size_pt", 11.0))
                        r.font.bold = meta.get("bold", False)
                        r.font.italic = meta.get("italic", False)
                        if meta.get("underline", False):
                            r.font.underline = True
                        is_pfx = any(tok.strip().lower() == pfx.lower() for pfx in field_prefixes)
                        if is_pfx:
                            r.font.color.rgb = RGBColor(0, 0, 0)
                        else:
                            r.font.color.rgb = RGBColor(*color)
                else:
                    run = p.add_run(full_text)
                    run.font.name = meta.get("font_name", "Times New Roman")
                    font_sz = meta.get("font_size_pt", 11.0)
                    run.font.size = Pt(font_sz)
                    run.font.bold = meta.get("bold", False)
                    run.font.italic = meta.get("italic", False)
                    if meta.get("underline", False):
                        run.font.underline = True
                    if color:
                        r, g, b = color
                        run.font.color.rgb = RGBColor(r, g, b)

            elif elem.element_type == "footer":
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(str(elem.content))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(7.5)
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def build_from_pages(self, pages: List[PageLayout], output_path: Union[str, Path]) -> Path:
        """
        Dựng toàn bộ tài liệu từ danh sách các PageLayout có phân trang chuẩn và hướng trang phù hợp.
        """
        # Chuẩn hóa chéo tên thực thể (co-reference resolution) giữa các trang
        all_canonical_names = []
        for p in pages:
            for elem in p.elements:
                if elem.element_type == "table" and getattr(elem.content, "extraction_method", "") == "signature_block_2col":
                    for c in elem.content.rows[-1]:
                        txt = c.text.strip()
                        if txt and len(txt.split()) >= 2:
                            all_canonical_names.append(txt)

        if all_canonical_names:
            import re
            from difflib import SequenceMatcher
            roles = ["Thủ kho", "Kế toán", "Giám đốc", "Phó giám đốc", "Trưởng phòng", "Nhân viên", "Đại diện"]
            for p in pages:
                for elem in p.elements:
                    if elem.element_type == "paragraph":
                        t_str = str(elem.content)
                        t_str = re.sub(r'\.{2,}', ' ', t_str)
                        t_str = re.sub(r'([a-zA-ZÀ-ỹ])\.\s+([a-zA-ZÀ-ỹ])', r'\1 \2', t_str)
                        words = t_str.split()
                        for c_name in all_canonical_names:
                            c_words = c_name.split()
                            n_w = len(c_words)
                            if len(words) >= n_w:
                                for i in range(len(words) - n_w + 1):
                                    window = ' '.join(words[i:i+n_w])
                                    if SequenceMatcher(None, window.lower(), c_name.lower()).ratio() >= 0.70:
                                        words = words[:i] + c_words + words[i+n_w:]
                                        break
                        t_str = ' '.join(words)
                        m_cv = re.search(r'chức\s*(?:vụ|vi)\s*:?', t_str, flags=re.I)
                        if m_cv:
                            cv_idx = m_cv.end()
                            suffix = t_str[cv_idx:].strip()
                            if suffix:
                                clean_suffix = re.sub(r'[^a-zA-ZÀ-ỹ\s]', ' ', suffix).strip()
                                s_words = clean_suffix.split()
                                for role in roles:
                                    w_cand = " ".join(s_words[:len(role.split())])
                                    if SequenceMatcher(None, w_cand.lower(), role.lower()).ratio() >= 0.40:
                                        t_str = t_str[:m_cv.start()].strip() + " Chức vụ: " + role
                                        break
                        elem.content = t_str

        doc = Document()
        for idx, page in enumerate(pages):
            is_landscape = (page.width > page.height)
            if idx == 0:
                sec = doc.sections[0]
            else:
                sec = doc.add_section(WD_SECTION_START.NEW_PAGE)

            if is_landscape:
                sec.orientation = WD_ORIENT.LANDSCAPE
                sec.page_width = Inches(self.page_height_in)  # 11.69"
                sec.page_height = Inches(self.page_width_in)  # 8.27"
            else:
                sec.orientation = WD_ORIENT.PORTRAIT
                sec.page_width = Inches(self.page_width_in)   # 8.27"
                sec.page_height = Inches(self.page_height_in) # 11.69"

            sec.top_margin = Inches(self.margin_in)
            sec.bottom_margin = Inches(self.margin_in)
            sec.left_margin = Inches(self.margin_in)
            sec.right_margin = Inches(self.margin_in)

            printable_w = sec.page_width.inches - sec.left_margin.inches - sec.right_margin.inches
            self.add_page_layout(doc, page, printable_width_in=printable_w)

        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

    def add_customs_reconciliation_table(self, document: Document, items: List[dict]) -> None:
        """
        Dựng Bảng Đối soát Kép Hải quan 7 cột (Dual-Audit Grid):
        1. Item No. (STT) - 5%
        2. Product Code (Mã hàng) - 22%
        3. Description & Specs (Mô tả & Thông số) - 28%
        4. H.S. Code (Mã số HS) - 10%
        5. Quantity & Unit (Số lượng & ĐVT) - 12%
        6. Invoice Reference (Số & Ngày HĐ) - 11%
        7. Audit Trail & Proof (Đối soát & Bằng chứng) - 12%
        """
        if not items:
            return

        col_pcts = [5.0, 22.0, 28.0, 10.0, 12.0, 11.0, 12.0]
        col_widths_in = [self.printable_width_in * (pct / 100.0) for pct in col_pcts]
        headers = [
            "Item No.", "Product Code", "Description & Specs",
            "H.S. Code", "Quantity & Unit", "Invoice Reference", "Audit Trail & Proof"
        ]

        table = document.add_table(rows=len(items) + 1, cols=7)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Thiết lập Header row: nền xám nhạt, lặp lại qua các trang, cantSplit
        header_row = table.rows[0]
        trPr = header_row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:tblHeader'))
        trPr.append(OxmlElement('w:cantSplit'))

        for c_idx, h_text in enumerate(headers):
            cell = header_row.cells[c_idx]
            cell.width = Inches(col_widths_in[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Tô màu nền xám nhẹ #F2F4F7 cho Header
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F2F4F7')
            tcPr.append(shd)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h_text)
            run.font.bold = True
            run.font.size = Pt(8.0)
            run.font.name = 'Calibri'
            set_cell_border(cell, top={'val': 'single', 'sz': 6, 'color': '888888'},
                                  bottom={'val': 'single', 'sz': 8, 'color': '444444'},
                                  left={'val': 'single', 'sz': 4, 'color': 'CCCCCC'},
                                  right={'val': 'single', 'sz': 4, 'color': 'CCCCCC'})

        # Điền dữ liệu các dòng
        for r_idx, item in enumerate(items):
            row = table.rows[r_idx + 1]
            # Chống gãy hàng qua trang
            r_trPr = row._tr.get_or_add_trPr()
            r_trPr.append(OxmlElement('w:cantSplit'))

            row_data = [
                str(item.get("item_no", r_idx + 1)),
                str(item.get("model_code", "")),
                str(item.get("description", "")),
                str(item.get("hs_code", "8525")),
                f"{item.get('quantity', '')} {item.get('unit', '')}".strip(),
                str(item.get("invoice_ref", "A24011321062\nMAY. 28, 2024")),
                f"{item.get('audit_proof', '✓ Verified')}\n[Raw: {item.get('raw_ocr', '')}]".strip()
            ]

            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = Inches(col_widths_in[c_idx])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05

                # Căn lề chuẩn công vụ
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif c_idx in [1, 2]:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif c_idx in [3, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif c_idx == 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:  # Audit trail
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                run = p.add_run(val)
                run.font.name = 'Calibri'
                if c_idx == 1:
                    run.font.bold = True
                    run.font.size = Pt(8.0)
                elif c_idx == 6:  # Audit trail font nhỏ xám
                    run.font.size = Pt(6.5)
                    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
                else:
                    run.font.size = Pt(8.0)

                # Viền kẻ bảng
                set_cell_border(cell, top={'val': 'single', 'sz': 4, 'color': 'E0E0E0'},
                                      bottom={'val': 'single', 'sz': 4, 'color': 'E0E0E0'},
                                      left={'val': 'single', 'sz': 4, 'color': 'E0E0E0'},
                                      right={'val': 'single', 'sz': 4, 'color': 'E0E0E0'})

    def build_customs_reconciliation_document(
        self,
        document_meta: dict,
        items: List[dict],
        output_path: Union[str, Path]
    ) -> Path:
        """
        Dựng toàn bộ Hồ sơ Đối soát Hải quan hoàn chỉnh gồm Header pháp lý và Bảng Đối soát Kép 7 cột.
        """
        doc = self.create_document()

        # Tiêu đề hồ sơ
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_after = Pt(2)
        r_title = title_p.add_run("BẢNG ĐỐI SOÁT HÀNG HÓA NHẬP KHẨU VÀ THÔNG QUAN HẢI QUAN")
        r_title.font.bold = True
        r_title.font.size = Pt(12.0)
        r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_p.paragraph_format.space_after = Pt(10)
        r_sub = sub_p.add_run("CUSTOMS RECONCILIATION & POST-CLEARANCE AUDIT REPORT")
        r_sub.font.bold = True
        r_sub.font.size = Pt(9.0)
        r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

        # Khối thông tin chứng từ (Metadata)
        meta_table = doc.add_table(rows=2, cols=2)
        meta_table.autofit = False
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        half_w = Inches(self.printable_width_in / 2.0)

        row0 = meta_table.rows[0]
        row0.cells[0].width = half_w
        row0.cells[1].width = half_w
        p00 = row0.cells[0].paragraphs[0]
        p00.add_run(f"Số Hồ sơ / Document Ref: ").font.bold = True
        p00.add_run(str(document_meta.get("doc_ref", "DS-K1107AM / HIKNHA0524-03")))

        p01 = row0.cells[1].paragraphs[0]
        p01.add_run(f"Số C/O / Certificate No: ").font.bold = True
        p01.add_run(str(document_meta.get("co_no", "24C3301M3571/00912 (CCPIT)")))

        row1 = meta_table.rows[1]
        row1.cells[0].width = half_w
        row1.cells[1].width = half_w
        p10 = row1.cells[0].paragraphs[0]
        p10.add_run(f"Nhà Xuất khẩu / Exporter: ").font.bold = True
        p10.add_run("HANGZHOU HIKVISION TECHNOLOGY CO., LTD.")

        p11 = row1.cells[1].paragraphs[0]
        p11.add_run(f"Người Nhập khẩu / Consignee: ").font.bold = True
        p11.add_run("CÔNG TY CỔ PHẦN NHÀ AN TOÀN")

        for r in meta_table.rows:
            r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c in r.cells:
                set_cell_border(c, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})

        sep_p = doc.add_paragraph()
        sep_p.paragraph_format.space_before = Pt(6)
        sep_p.paragraph_format.space_after = Pt(6)

        # Chèn Bảng 7 Cột
        self.add_customs_reconciliation_table(doc, items)

        # Chân trang ký duyệt
        sign_p = doc.add_paragraph()
        sign_p.paragraph_format.space_before = Pt(16)
        sign_table = doc.add_table(rows=1, cols=2)
        sign_table.autofit = False
        sign_table.rows[0].cells[0].width = half_w
        sign_table.rows[0].cells[1].width = half_w

        p_s0 = sign_table.rows[0].cells[0].paragraphs[0]
        p_s0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s0.add_run("NGƯỜI LẬP BIỂU / OPERATOR\n").font.bold = True
        p_s0.add_run("(Ký và ghi rõ họ tên)")

        p_s1 = sign_table.rows[0].cells[1].paragraphs[0]
        p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s1.add_run("CÁN BỘ ĐỐI SOÁT / CUSTOMS SPECIALIST\n").font.bold = True
        p_s1.add_run("(Ký, đóng dấu xác nhận)")

        for c in sign_table.rows[0].cells:
            set_cell_border(c, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})

        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

    def build(self, table_results: List[TableResult], output_path: Union[str, Path]) -> Path:
        """Dựng tài liệu trực tiếp từ danh sách TableResult (tương thích ngược cho unit test)."""
        doc = self.create_document()
        for idx, t in enumerate(table_results):
            if idx > 0:
                doc.add_page_break()
            self.add_table(doc, t)
        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

